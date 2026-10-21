"""Documentos em Word (.docx): texto, tabelas, cabeçalho e imagens coladas viram blocos para a IA."""
import base64
import io
import sys
from pathlib import Path

import pytest
from docx import Document
from fastapi import HTTPException

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import main  # noqa: E402

# PNG 1x1 válido
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")


def _docx() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "FICHA DE ADMISSÃO"
    doc.add_paragraph("Nome: Maria da Silva")
    tabela = doc.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text, tabela.cell(0, 1).text = "CPF", "123.456.789-09"
    tabela.cell(1, 0).text, tabela.cell(1, 1).text = "Nascimento", "01/02/1990"
    doc.add_picture(io.BytesIO(PNG))
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def test_docx_vira_texto_e_imagens():
    blocos = main._blocos_docx("ficha.docx", _docx())
    texto = blocos[0]["text"]
    assert "FICHA DE ADMISSÃO" in texto and "Nome: Maria da Silva" in texto and "CPF | 123.456.789-09" in texto
    imagens = [b for b in blocos if b["type"] == "image"]
    assert len(imagens) == 1 and imagens[0]["source"]["media_type"] == "image/png"


def test_reconhece_docx_por_tipo_ou_extensao():
    assert main._eh_docx(main.DOCX_MEDIA_TYPE, "x")
    assert main._eh_docx("application/octet-stream", "Ficha.DOCX")
    assert not main._eh_docx("application/pdf", "a.pdf")


def test_arquivo_que_nao_e_docx_da_erro_claro():
    with pytest.raises(HTTPException) as erro:
        main._blocos_docx("antigo.docx", b"isto nao e um zip")
    assert erro.value.status_code == 400 and ".doc antigo" in erro.value.detail
